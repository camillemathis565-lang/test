# Export the "densité d équipements" chart data to Excel. Run from cross_border/: python report_docx/export_density_xlsx.py
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

d = pd.read_csv('report_docx/data/communes_retail_profile.csv', dtype={'code': str})
d['retail'] = d.small_food + d.large_food + d.non_food
d['zone'] = d.dist_km < 20
EQ = [("Super et hypermarchés", "large_food"), ("Commerces de détail", "retail"), ("Stations-service", "fuel"), ("Buralistes", "tobacconists")]

wb = Workbook()
H = Font(bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="0070C0")
thin = Side(style="thin", color="D5DCE4"); B = Border(top=thin, bottom=thin, left=thin, right=thin)
def header(ws, row, cols):
    for j, c in enumerate(cols, 1):
        x = ws.cell(row=row, column=j, value=c); x.font = H; x.fill = HF; x.alignment = Alignment(wrap_text=True, vertical="center"); x.border = B
def widths(ws, w):
    for i, v in enumerate(w, 1): ws.column_dimensions[get_column_letter(i)].width = v

# --- Sheet 1: chart data
ws = wb.active; ws.title = "Graphique"
ws["A1"] = "Densité d'équipements, zone frontalière (0–20 km) par rapport à la France métropolitaine (écart en %)"; ws["A1"].font = Font(bold=True, size=12)
ws["A2"] = "Densité = nombre d'équipements / population × 10 000. Écart = densité zone / densité France − 1. Formules visibles dans les colonnes E, H et I."
header(ws, 4, ["Barre du graphique", "Zone", "Équipements (nombre)", "Population 2023", "Densité (pour 10 000 hab.)", "Équipements France métro.", "Population France métro.", "Densité France métro.", "Écart (%)"])
nat = {c: d[c].sum() for _, c in EQ}; natpop = d['pop'].sum()
rows = []
for lab, c in EQ:
    z = d[d.zone]; rows.append((lab, "Zone frontalière (toutes frontières)", z[c].sum(), z['pop'].sum(), c))
for lab, c in [("Stations-service, frontière luxembourgeoise", "fuel"), ("Buralistes, frontière luxembourgeoise", "tobacconists")]:
    z = d[d.zone & (d.neighbour == "Luxembourg")]; rows.append((lab, "Zone frontalière, frontière luxembourgeoise", z[c].sum(), z['pop'].sum(), c))
for i, (lab, zone, n, p, c) in enumerate(rows, 5):
    ws.cell(i, 1, lab); ws.cell(i, 2, zone); ws.cell(i, 3, int(n)); ws.cell(i, 4, int(p))
    ws.cell(i, 5, f"=C{i}/D{i}*10000"); ws.cell(i, 6, int(nat[c])); ws.cell(i, 7, int(natpop))
    ws.cell(i, 8, f"=F{i}/G{i}*10000"); ws.cell(i, 9, f"=E{i}/H{i}-1")
    for j in range(1, 10): ws.cell(i, j).border = B
    for j in (3, 4, 6, 7): ws.cell(i, j).number_format = "# ##0"
    for j in (5, 8): ws.cell(i, j).number_format = "0.00"
    ws.cell(i, 9).number_format = "+0%;-0%"
widths(ws, [42, 40, 14, 14, 14, 16, 16, 14, 11]); ws.row_dimensions[4].height = 45

# --- Sheet 2: by border
ws = wb.create_sheet("Par frontière")
ws["A1"] = "Zone frontalière (communes à moins de 20 km), par pays voisin le plus proche"; ws["A1"].font = Font(bold=True, size=12)
cols = ["Frontière", "Communes", "Population 2023"] + [f"{l} (nombre)" for l, _ in EQ] + [f"{l} pour 10 000 hab." for l, _ in EQ] + [f"{l} : écart vs France" for l, _ in EQ]
header(ws, 3, cols)
groups = [(nb, d[d.zone & (d.neighbour == nb)]) for nb in ["Belgium", "Luxembourg", "Germany", "Switzerland", "Italy", "Spain"]] + [("Toutes frontières", d[d.zone]), ("France métropolitaine (hors Corse)", d)]
FR = {"Belgium": "Belgique", "Luxembourg": "Luxembourg", "Germany": "Allemagne", "Switzerland": "Suisse", "Italy": "Italie (et Monaco)", "Spain": "Espagne (et Andorre)"}
last = 3 + len(groups)
for i, (nb, g) in enumerate(groups, 4):
    ws.cell(i, 1, FR.get(nb, nb)); ws.cell(i, 2, len(g)); ws.cell(i, 3, int(g['pop'].sum()))
    for k, (_, c) in enumerate(EQ):
        cn = 4 + k; ws.cell(i, cn, int(g[c].sum()))
        L = get_column_letter(cn); dc = 8 + k; ws.cell(i, dc, f"={L}{i}/C{i}*10000").number_format = "0.00"
        D = get_column_letter(dc); ws.cell(i, 12 + k, f"={D}{i}/{D}${last}-1").number_format = "+0%;-0%"
    for j in range(1, 16): ws.cell(i, j).border = B
    ws.cell(i, 3).number_format = "# ##0"
    if i >= last - 1:
        for j in range(1, 16): ws.cell(i, j).font = Font(bold=True)
widths(ws, [30, 10, 14] + [13] * 12); ws.row_dimensions[3].height = 60

# --- Sheet 3: communes
ws = wb.create_sheet("Communes")
c = d[['code', 'LAU_NAME', 'neighbour', 'dist_km', 'zone', 'pop', 'large_food', 'retail', 'fuel', 'tobacconists']].copy()
c['neighbour'] = c.neighbour.map(FR); c['dist_km'] = c.dist_km.round(1); c['zone'] = c.zone.map({True: "Oui", False: "Non"})
header(ws, 1, ["Code commune", "Commune", "Frontière la plus proche", "Distance à la frontière (km)", "Zone frontalière (< 20 km)", "Population 2023", "Super et hypermarchés", "Commerces de détail", "Stations-service", "Buralistes (2018)"])
for r in c.itertuples(index=False): ws.append(list(r))
ws.auto_filter.ref = f"A1:J{len(c)+1}"; ws.freeze_panes = "A2"
widths(ws, [12, 32, 22, 14, 14, 14, 12, 12, 12, 12])

# --- Sheet 4: method
ws = wb.create_sheet("Méthode et sources")
txt = [
 ("Méthode", True),
 ("Champ : communes de France métropolitaine hors Corse (découpage Eurostat GISCO, LAU 2024).", False),
 ("Distance : distance à vol d'oiseau entre un point intérieur de la commune (point représentatif) et la frontière terrestre la plus proche. Zone frontalière = communes à moins de 20 km.", False),
 ("Frontière la plus proche : Monaco regroupé avec l'Italie, Andorre avec l'Espagne.", False),
 ("Densité = nombre d'équipements / population municipale 2023 × 10 000, calculée sur l'ensemble des communes de la zone (et non en moyenne des densités communales).", False),
 ("Écart = densité de la zone / densité de la France métropolitaine − 1. La France métropolitaine inclut la zone frontalière.", False),
 ("", False),
 ("Définitions (BPE 2025)", True),
 ("Super et hypermarchés : B104 (hypermarché et grand magasin, ≥ 2 500 m²) et B105 (supermarché et magasin multi-commerce, 400 à 2 500 m²).", False),
 ("Commerces de détail : commerces alimentaires (B201, B202, B204 à B210), super et hypermarchés (B104, B105) et commerces non alimentaires spécialisés (B103, B302 à B304, B306 à B313, B315, B317 à B319, B321 à B325). Stations-service exclues.", False),
 ("Stations-service : B316. Attention : la BPE ne couvre que les stations ayant vendu au moins 500 000 litres l'année précédente ; les plus petites n'y figurent que sur la base du volontariat.", False),
 ("Buralistes : annuaire DGDDI 2018, géolocalisé ; rapporté à la population 2023.", False),
 ("", False),
 ("Limites", True),
 ("Lecture descriptive : une densité plus faible est cohérente avec les achats à l'étranger mais peut aussi refléter la structure du peuplement ou le tourisme.", False),
 ("La distance à vol d'oiseau ignore le relief et les points de passage.", False),
 ("", False),
 ("Sources", True),
 ("Insee (2026), Base permanente des équipements 2025, https://www.insee.fr/fr/statistiques/8217525", False),
 ("Insee (2025), Populations de référence 2023, https://api.insee.fr/melodi/catalog/DS_POPULATIONS_REFERENCE", False),
 ("DGDDI (2018), Annuaire des buralistes de France métropolitaine 2018, https://data.economie.gouv.fr/explore/dataset/annuaire-des-buralistes-de-france-metropolitaine-2018/", False),
 ("Eurostat (2024), Local Administrative Units (LAU) 2024 et frontières des pays, GISCO, https://gisco-services.ec.europa.eu/distribution/v2/lau/", False),
 ("Calculs : cross_border/report_docx/retail_profile.py", False),
]
for i, (t, b) in enumerate(txt, 1):
    x = ws.cell(i, 1, t); x.font = Font(bold=b, size=12 if b else 11); x.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 130
wb.save('report_docx/Densite_equipements_zone_frontaliere.xlsx')
print("ok")
